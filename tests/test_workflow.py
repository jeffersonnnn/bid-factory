import io
import json
import sqlite3
import zipfile
from datetime import datetime, timezone
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from app.db import transaction
from conftest import post, upload, base, prepared, ready, state, REVIEW


def test_complete_package_contains_grounded_claims_and_real_documents(client):
    co,bid,_=ready(client)
    assert state(client,bid)['ready'] is True
    r=client.post(f'/api/bids/{bid}/export',json={'final':True})
    assert r.status_code==200,r.text
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        assert {'Proposal.docx','Proposal.pdf','internal/Provenance.json','internal/Compliance-Matrix.xlsx','internal/Submission-Checklist.pdf'}<=set(z.namelist())
        manifest=json.loads(z.read('internal/Provenance.json'))
        assert manifest['final'] is True and manifest['page_count']<=25
        assert any(n.startswith('original-attachments/') for n in z.namelist())
        assert 'ABCD1234EFGH' in ''.join(p.extract_text() for p in PdfReader(io.BytesIO(z.read('Proposal.pdf'))).pages)
        doc=Document(io.BytesIO(z.read('Proposal.docx')))
        text='\n'.join(p.text for p in doc.paragraphs)
        for section in manifest['claims']:
            for claim in section['claims']:
                assert claim['text'] in text
                assert claim['evidence_id']
        book=load_workbook(io.BytesIO(z.read('internal/Compliance-Matrix.xlsx')))
        assert book.active.max_row==9


def test_missing_uei_is_blocker_not_reminder(client):
    _,bid,_=base(client)
    data=state(client,bid)
    assert any(f['rule_id']=='RULE-010' for f in data['blockers'])
    assert not data['ready']
    assert client.post(f'/api/bids/{bid}/export',json={'final':True}).status_code==409


def test_missing_eligibility_is_no_go(client):
    _,bid,_=base(client)
    assert post(client,f'/bids/{bid}/qualify')['status']=='NO-GO'


def test_explicit_wrong_eligibility_is_no_go(client):
    co,bid,_=prepared(client)
    req=next(r for r in state(client,bid)['requirements'] if r['kind']=='eligibility')
    answer=post(client,f'/bids/{bid}/answers',{'requirement_id':req['id'],'text':'Example Services is not a small business.',
        'customer_name':'Test customer','confirmed':True,'fact_key':'negative_eligibility'})
    post(client,f"/requirements/{req['id']}/mapping",{**REVIEW,'evidence_ids':[answer['evidence_id']]})
    assert state(client,bid)['qualification']['status']=='NO-GO'


def test_pricing_document_cannot_replace_customer_input(client):
    co,bid,_=base(client)
    doc=upload(client,f'/companies/{co}/documents','price-sheet.txt')['id']
    post(client,f'/companies/{co}/extract')
    evidence=client.get('/api/companies/'+co).json()['evidence'][0]
    post(client,f"/evidence/{evidence['id']}/review",{**REVIEW,'approved':True})
    req=next(r for r in state(client,bid)['requirements'] if r['kind']=='pricing')
    post(client,f"/requirements/{req['id']}/mapping",{**REVIEW,'evidence_ids':[evidence['id']]})
    data=state(client,bid)
    assert any(f['rule_id']=='RULE-009' for f in data['blockers'])
    assert any('Customer input required' in q['text'] for q in data['questions'])


def test_invented_quote_is_rejected(client):
    co,bid,doc=base(client)
    r=client.post(f'/api/bids/{bid}/requirements',json={'document_id':doc,'page':1,'quote':'We have 12 years of experience.',
        'text':'12 years','kind':'technical','key':'experience'})
    assert r.status_code==422


def test_prompt_injection_is_not_evidence(client):
    co=post(client,'/companies',{'name':'Hostile test'})['id']
    upload(client,f'/companies/{co}/documents','hostile-company.txt')
    post(client,f'/companies/{co}/extract')
    evidence=client.get('/api/companies/'+co).json()['evidence']
    assert len(evidence)==2
    assert all('50 federal contracts' not in e['statement'] for e in evidence)


def test_conflicting_experience_requires_customer_resolution(client):
    co,bid,_=base(client)
    upload(client,f'/companies/{co}/documents','hostile-company.txt')
    post(client,f'/companies/{co}/extract')
    data=state(client,bid)
    assert data['conflicts']
    assert any('Conflicting evidence' in f['finding'] for f in data['blockers'])


def test_amendment_keeps_25_to_20_page_lineage(client):
    co,bid,doc=base(client)
    old=next(r for r in state(client,bid)['requirements'] if r['kind']=='page_limit')
    upload(client,f'/bids/{bid}/documents','amendment-0001.txt',kind='amendment',sequence=1)
    post(client,f'/bids/{bid}/extract')
    data=state(client,bid)
    current=next(r for r in data['requirements'] if r['kind']=='page_limit')
    assert json.loads(current['constraints'])['max_pages']==20
    assert current['supersedes']==old['id']
    assert any(r['id']==old['id'] for r in data['ledger'])
    assert any(f['rule_id']=='RULE-008' for f in data['blockers'])


def test_old_amendment_cannot_override_later_amendment(client):
    _,bid,_=base(client)
    upload(client,f'/bids/{bid}/documents','amendment-0002.txt',b'Maximum proposal length revised to 18 pages.',kind='amendment',sequence=2)
    post(client,f'/bids/{bid}/extract')
    latest=next(r for r in state(client,bid)['requirements'] if r['kind']=='page_limit')
    olddoc=upload(client,f'/bids/{bid}/documents','amendment-0001.txt',kind='amendment',sequence=1)['id']
    r=client.post(f'/api/bids/{bid}/requirements',json={'document_id':olddoc,'page':1,'quote':'Maximum proposal length revised to 20 pages.',
        'text':'20 pages','kind':'page_limit','key':latest['key'],'supersedes':latest['id'],'supersession_reason':'Wrong order test'})
    assert r.status_code==422


def test_duplicate_upload_and_extraction_are_idempotent(client):
    _,bid,doc=base(client)
    n=len(state(client,bid)['ledger'])
    assert upload(client,f'/bids/{bid}/documents','simple-rfq.txt',kind='solicitation')['duplicate']
    post(client,f'/bids/{bid}/extract')
    assert len(state(client,bid)['ledger'])==n


def test_compound_requirement_becomes_four_atoms(client):
    co=post(client,'/companies',{'name':'Compound test'})['id']
    bid=post(client,'/bids',{'company_id':co,'title':'Compound RFQ'})['id']
    upload(client,f'/bids/{bid}/documents','compound.txt',b'Offeror must provide technical approach, three references, UEI and price sheet.',kind='solicitation')
    post(client,f'/bids/{bid}/extract')
    reqs=state(client,bid)['requirements']
    assert len(reqs)==4
    assert {r['kind'] for r in reqs}=={'technical','past_performance','administrative','pricing'}


def test_performance_scope_is_not_submission_scope(client):
    _,bid,_=base(client)
    upload(client,f'/bids/{bid}/documents','pws.txt',b'Contractor shall clean bathrooms daily.',kind='solicitation')
    post(client,f'/bids/{bid}/extract')
    assert next(r for r in state(client,bid)['requirements'] if 'bathrooms' in r['text'])['scope']=='performance'


def test_cross_company_evidence_is_rejected(client):
    co,bid,_=prepared(client)
    other=post(client,'/companies',{'name':'Other company'})['id']
    upload(client,f'/companies/{other}/documents','company.txt')
    post(client,f'/companies/{other}/extract')
    eid=client.get('/api/companies/'+other).json()['evidence'][0]['id']
    req=state(client,bid)['requirements'][0]
    assert client.post(f"/api/requirements/{req['id']}/mapping",json={**REVIEW,'evidence_ids':[eid]}).status_code==422


def test_later_upload_revokes_ready_and_final_export(client):
    _,bid,_=ready(client)
    upload(client,f'/bids/{bid}/documents','amendment-0001.txt',kind='amendment',sequence=1)
    assert state(client,bid)['ready'] is False
    assert client.post(f'/api/bids/{bid}/export',json={'final':True}).status_code==409


def test_claim_tampering_cannot_pass_audit(client):
    _,bid,_=prepared(client)
    post(client,f'/bids/{bid}/generate')
    with transaction() as db:
        db.execute("UPDATE proposal_claims SET text='We have 50 federal contracts.' WHERE id=(SELECT id FROM proposal_claims LIMIT 1)")
    data=state(client,bid)
    assert any(f['rule_id']=='RULE-002' for f in data['blockers'])
    assert client.post(f'/api/bids/{bid}/approve',json=REVIEW).status_code==409


def test_mandatory_requirement_cannot_be_deleted(client):
    _,bid,_=base(client)
    with transaction() as db:
        try:
            db.execute('DELETE FROM requirements WHERE bid_id=?',(bid,))
            assert False,'Deletion unexpectedly succeeded'
        except sqlite3.IntegrityError:
            pass


def test_missing_attachment_blocks_even_with_customer_price(client):
    _,bid,_=base(client)
    req=next(r for r in state(client,bid)['requirements'] if r['kind']=='pricing')
    answer=post(client,f'/bids/{bid}/answers',{'requirement_id':req['id'],'text':'The price is $1200 per month.',
        'customer_name':'Customer Test','confirmed':True})
    post(client,f"/requirements/{req['id']}/mapping",{**REVIEW,'evidence_ids':[answer['evidence_id']]})
    assert next(r for r in state(client,bid)['requirements'] if r['id']==req['id'])['status']=='PARTIAL'


def test_past_performance_requires_three_distinct_references(client):
    _,bid,_=prepared(client)
    data=state(client,bid)
    req=next(r for r in data['requirements'] if r['kind']=='past_performance')
    ids=[e['id'] for e in data['evidence'] if e['kind']=='past_performance'][:2]
    post(client,f"/requirements/{req['id']}/mapping",{**REVIEW,'evidence_ids':ids})
    data=state(client,bid)
    assert next(r for r in data['requirements'] if r['id']==req['id'])['status']=='PARTIAL'
    assert any('1 more distinct' in q['text'] for q in data['questions'])


def test_fake_certification_cannot_be_claimed_without_evidence(client):
    _,bid,_=base(client)
    upload(client,f'/bids/{bid}/documents','cert.txt',b'Offeror must provide a current ISO 9001 certification.',kind='solicitation')
    post(client,f'/bids/{bid}/extract')
    assert state(client,bid)['qualification']['status']=='NO-GO'


def test_unreviewed_extraction_blocks_drafting(client):
    _,bid,_=base(client)
    assert client.post(f'/api/bids/{bid}/generate',json={}).status_code==409


def test_red_team_finding_survives_rerun(client):
    _,bid,_=prepared(client)
    post(client,f'/bids/{bid}/generate')
    first=post(client,f'/bids/{bid}/audit',{**REVIEW,'mode':'human','independent_review_complete':True,'findings':['Missing contract number for reference 1.']})
    assert first['status']=='FAILED'
    again=post(client,f'/bids/{bid}/audit',{**REVIEW,'mode':'human','independent_review_complete':True})
    assert again['status']=='FAILED'
    f=next(f for f in state(client,bid)['audit_findings'] if f['rule_id']=='RED-TEAM')
    assert client.post(f"/api/findings/{f['id']}/resolve",json=REVIEW).status_code==409


def test_csrf_and_untrusted_hosts_are_rejected(client):
    assert client.post('/api/companies',json={'name':'Attack'},headers={'origin':'https://attacker.example'}).status_code==403
    assert client.get('/api/workspace',headers={'host':'attacker.example'}).status_code==400


def test_production_password_protects_data(client,monkeypatch):
    monkeypatch.setenv('APP_PASSWORD','a-long-test-password')
    assert client.get('/api/workspace').status_code==401
    assert client.get('/api/workspace',auth=('operator','a-long-test-password')).status_code==200


def test_invalid_files_fail_without_partial_records(client):
    co=post(client,'/companies',{'name':'File test'})['id']
    r=client.post(f'/api/companies/{co}/documents',files={'file':('fake.pdf',b'not a pdf')})
    assert r.status_code==422
    assert client.get('/api/companies/'+co).json()['documents']==[]
