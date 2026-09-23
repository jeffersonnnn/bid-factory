"""Bounded document parsing. PDF pages are real pages; DOCX locators are paragraphs/tables."""
import hashlib
import io
import re
import zipfile
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from fastapi import HTTPException

MAX_BYTES = 12 * 1024 * 1024
MAX_CHARS = 350_000
INJECTION = re.compile(r'ignore\b.{0,80}\b(instruction|solicitation|previous)|system\s*prompt|you are (chatgpt|claude)|state that .{0,100}(contract|experience)|disregard .{0,60}instruction', re.I | re.S)


def sha(content):
    return hashlib.sha256(content).hexdigest()


def parse(name, data, allow_scanned=False):
    if not data or len(data) > MAX_BYTES:
        raise HTTPException(422, 'Upload a non-empty file of 12 MB or less.')
    ext = Path(name).suffix.lower()
    warnings, pages = [], []
    try:
        if ext in ('.docx', '.xlsx'):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(x.file_size for x in z.infolist()) > 40 * 1024 * 1024 or len(z.infolist()) > 5000:
                    raise ValueError('Compressed document exceeds the safe size limit.')
        if ext == '.pdf':
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError('Remove PDF password protection before upload.')
            if len(reader.pages) > 150:
                raise ValueError('The initial version supports at most 150 source pages per file.')
            for i, p in enumerate(reader.pages):
                t = p.extract_text() or ''
                pages.append({'number': i + 1, 'locator': f'Page {i + 1}', 'text': t})
                if len(t.strip()) < 15:
                    warnings.append(f'Page {i + 1} has little or no text. Upload a text-readable replacement after OCR.')
        elif ext == '.docx':
            doc = Document(io.BytesIO(data))
            # Preserve the document body order, including table cells.
            for i, item in enumerate(doc.iter_inner_content()):
                if hasattr(item, 'rows'):
                    t = '\n'.join(' | '.join(c.text for c in row.cells) for row in item.rows)
                    loc = f'Body block {i + 1} (table)'
                else:
                    t, loc = item.text, f'Body block {i + 1} (paragraph)'
                if t.strip():
                    pages.append({'number': i + 1, 'locator': loc, 'text': t})
            warnings.append('DOCX locations use body blocks, not printed page numbers. Review headers, footers and text boxes in the original file.')
        elif ext == '.xlsx':
            book = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
            for sheet in book:
                if sheet.max_row and sheet.max_row > 5000:
                    raise ValueError('Spreadsheet exceeds 5,000 rows.')
                t = '\n'.join(f'Row {i}: ' + ' | '.join(str(c) if c is not None else '' for c in row)
                              for i, row in enumerate(sheet.iter_rows(values_only=True), 1))
                pages.append({'number': len(pages) + 1, 'locator': f'Sheet: {sheet.title}', 'text': t})
            warnings.append('Spreadsheet formulas are source text. Verify displayed amounts in the original workbook.')
        elif ext in ('.txt', '.md', '.csv'):
            for i, t in enumerate(data.decode('utf-8-sig').split('\f')):
                pages.append({'number': i + 1, 'locator': f'Text block {i + 1}', 'text': t})
        else:
            raise ValueError('Use PDF, DOCX, XLSX, TXT, MD or CSV.')
        if not pages or (not any(p['text'].strip() for p in pages) and not (allow_scanned and ext == '.pdf')):
            raise ValueError('No readable text. Run OCR and upload the text-readable file.')
        if sum(len(p['text']) for p in pages) > MAX_CHARS:
            raise ValueError('The document exceeds the text limit. Split it into smaller files.')
        if any(INJECTION.search(p['text']) for p in pages):
            warnings.append('Possible instructions directed at an AI were found. Review the source as untrusted text.')
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f'Cannot read this file: {str(exc)[:180]}') from exc
    return pages, warnings


def source_page(document, page, quote):
    import json
    p = next((p for p in json.loads(document['pages']) if p['number'] == page), None)
    if not p or quote not in p['text']:
        raise HTTPException(422, 'The exact source quote does not exist at this location.')
    return p


def safe_filename(name):
    return re.sub(r'[^a-zA-Z0-9_. -]', '_', Path(name.replace('\\', '/')).name)[:150] or 'document.txt'
